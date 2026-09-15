"""Adapter and shared helpers for one Uber expense source workbook."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from deal_expenses.models import SourceSummary, ValidationResult


REQUIRED_HEADERS = frozenset(
    {
        "Request Date (UTC)",
        "First Name",
        "Last Name",
        "Employee ID",
        "Transaction Amount USD",
    }
)


def normalize_uber_header(value: Any) -> str:
    """Return a case-insensitive, whitespace-normalized header."""
    return " ".join(str(value or "").strip().split()).casefold()


def decimal_amount(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as error:
        raise ValueError(f"Transaction Amount USD is not numeric: {value!r}") from error


def worksheet_headers(worksheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    duplicates: list[str] = []
    for column, cell in enumerate(worksheet[1], start=1):
        header = normalize_uber_header(cell.value)
        if not header:
            continue
        if header in headers:
            duplicates.append(str(cell.value))
            continue
        headers[header] = column
    if duplicates:
        raise ValueError(f"Duplicate source headers in '{worksheet.title}': {duplicates}")
    return headers


class UberSourceAdapter:
    """Read the one worksheet in an Uber workbook with the required headers."""

    key = "uber"
    display_name = "Uber report"

    def select_worksheet(self, workbook):
        required_headers = {normalize_uber_header(header) for header in REQUIRED_HEADERS}
        matches = []
        for worksheet in workbook.worksheets:
            try:
                headers = worksheet_headers(worksheet)
            except ValueError:
                continue
            if required_headers.issubset(headers):
                matches.append(worksheet)
        if len(matches) != 1:
            names = ", ".join(worksheet.title for worksheet in matches) or "none"
            raise ValueError(f"Expected exactly one Uber source worksheet with required headers; found {names}.")
        return matches[0]

    def validate(self, source_path: Path) -> list[ValidationResult]:
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=False)
            try:
                worksheet = self.select_worksheet(workbook)
                headers = worksheet_headers(worksheet)
            finally:
                workbook.close()
        except Exception as error:
            return [ValidationResult(self.display_name, False, str(error))]

        missing = [header for header in REQUIRED_HEADERS if normalize_uber_header(header) not in headers]
        if missing:
            return [ValidationResult(self.display_name, False, "Missing required column(s): " + ", ".join(missing))]
        return [ValidationResult(self.display_name, True, f"Worksheet '{worksheet.title}' is valid.")]

    def summarize(self, source_path: Path, year: int) -> SourceSummary:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            worksheet = self.select_worksheet(workbook)
            headers = worksheet_headers(worksheet)
            amount_column = headers[normalize_uber_header("Transaction Amount USD")]
            rows = 0
            total = Decimal("0")
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if all(value in (None, "") for value in row):
                    continue
                rows += 1
                total += decimal_amount(row[amount_column - 1])
        finally:
            workbook.close()

        if not rows:
            raise ValueError(f"{self.display_name} has no data rows.")
        return SourceSummary(self.key, self.display_name, rows, float(total))
