"""Build the latest-record HR refresh plan used by the Uber workflow."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HR_SHEET_PREFIX = "HR HC Combined"
HR_SOURCE_SHEET = "To be Copy Pasted from HR"
NETWORK_ID_HEADER = "Network ID"
EFFECTIVE_DATE_HEADER = "Effective Date"
DUPLICATE_HEADERS = frozenset({"duplicate", "duplicates", "duplicate?"})
REQUIRED_HR_HEADERS = (NETWORK_ID_HEADER, EFFECTIVE_DATE_HEADER, "Name", "Cost Center Name")


def normalize_hr_value(value: Any) -> str:
    """Normalize a header, name, or identifier for case-insensitive matching."""
    return " ".join(str(value or "").strip().split()).casefold()


def normalize_network_id(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    value = str(value).strip()
    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]
    return value.casefold()


def coerce_excel_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    value = str(value or "").strip()
    for pattern in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            continue
    return None


def _headers(worksheet, row_number: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    duplicates: list[str] = []
    for column, cell in enumerate(worksheet[row_number], start=1):
        header = normalize_hr_value(cell.value)
        if not header:
            continue
        if header in headers:
            duplicates.append(str(cell.value))
        headers[header] = column
    if duplicates:
        raise ValueError(f"Duplicate nonblank headers in '{worksheet.title}': {duplicates}")
    return headers


def _header_row(worksheet, required_header: str, maximum_rows: int = 20) -> int:
    header = normalize_hr_value(required_header)
    for row_number in range(1, min(maximum_rows, worksheet.max_row) + 1):
        values = {normalize_hr_value(cell.value) for cell in worksheet[row_number]}
        if header in values:
            return row_number
    raise ValueError(f"Could not find header '{required_header}' in the first {maximum_rows} rows of '{worksheet.title}'.")


def _master_worksheet(workbook):
    matches = [worksheet for worksheet in workbook.worksheets if worksheet.title.startswith(HR_SHEET_PREFIX)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one worksheet beginning '{HR_SHEET_PREFIX}'; found {[sheet.title for sheet in matches]}.")
    return matches[0]


def _source_worksheet(workbook):
    if HR_SOURCE_SHEET in workbook.sheetnames:
        return workbook[HR_SOURCE_SHEET]
    matches = []
    for worksheet in workbook.worksheets:
        try:
            row_number = _header_row(worksheet, NETWORK_ID_HEADER)
            headers = _headers(worksheet, row_number)
        except ValueError:
            continue
        if normalize_hr_value(EFFECTIVE_DATE_HEADER) in headers:
            matches.append(worksheet)
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one Headcount source worksheet; found {[sheet.title for sheet in matches]}.")
    return matches[0]


def _required_column(headers: dict[str, int], header: str, worksheet_name: str) -> int:
    column = headers.get(normalize_hr_value(header))
    if column is None:
        raise ValueError(f"Required HR column '{header}' is missing from '{worksheet_name}'.")
    return column


def _candidate_rows(worksheet, header_row: int, headers: dict[str, int], origin: str, output_headers: dict[str, int]) -> list[dict[str, Any]]:
    network_column = _required_column(headers, NETWORK_ID_HEADER, worksheet.title)
    effective_date_column = _required_column(headers, EFFECTIVE_DATE_HEADER, worksheet.title)
    final_column = max(headers.values())
    candidates: list[dict[str, Any]] = []
    for row_number, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, max_col=final_column, values_only=True), start=header_row + 1):
        if all(value in (None, "") for value in values):
            continue
        network_id = normalize_network_id(values[network_column - 1])
        effective_date = coerce_excel_date(values[effective_date_column - 1])
        if network_id and effective_date is None:
            raise ValueError(f"{origin.title()} HR row {row_number} has Network ID {values[network_column - 1]!r} but no valid Effective Date.")
        if origin == "master":
            output_values = tuple(values[index - 1] if index <= len(values) else None for index in range(1, max(output_headers.values()) + 1))
        else:
            output = [None] * max(output_headers.values())
            for header, source_column in headers.items():
                output[output_headers[header] - 1] = values[source_column - 1]
            output_values = tuple(output)
        candidates.append({"origin": origin, "row_number": row_number, "network_id": network_id, "effective_date": effective_date, "values": output_values})
    return candidates


def _retain_latest(candidates: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    by_network_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    retained = []
    for row in candidates:
        if row["network_id"]:
            by_network_id[row["network_id"]].append(row)
        else:
            retained.append(row)
    removed = 0
    ties: list[str] = []
    for network_id, rows in by_network_id.items():
        latest_date = max(row["effective_date"] for row in rows)
        latest_rows = [row for row in rows if row["effective_date"] == latest_date]
        if len(latest_rows) != 1:
            ties.append(network_id)
            continue
        retained.append(latest_rows[0])
        removed += len(rows) - 1
    if ties:
        raise ValueError("Cannot choose a latest record for Network ID(s) with equal maximum Effective Date: " + ", ".join(ties[:20]))
    retained.sort(key=lambda row: (row["origin"] != "master", row["row_number"]))
    return retained, removed


@dataclass(frozen=True)
class HrRefreshPlan:
    master_sheet_name: str
    master_header_row: int
    master_headers: dict[str, int]
    retained_rows: list[tuple[Any, ...]]
    existing_rows: int
    source_rows: int
    superseded_rows: int
    duplicate_column: int | None


def build_hr_refresh_plan(master_path: Path, source_path: Path) -> HrRefreshPlan:
    """Merge master and Headcount HR rows, retaining the newest record per Network ID."""
    master_workbook = load_workbook(master_path, read_only=True, data_only=False)
    source_workbook = load_workbook(source_path, read_only=True, data_only=False)
    try:
        master_sheet = _master_worksheet(master_workbook)
        source_sheet = _source_worksheet(source_workbook)
        master_header_row = _header_row(master_sheet, NETWORK_ID_HEADER)
        source_header_row = _header_row(source_sheet, NETWORK_ID_HEADER)
        master_headers = _headers(master_sheet, master_header_row)
        source_headers = _headers(source_sheet, source_header_row)
        for header in REQUIRED_HR_HEADERS:
            _required_column(master_headers, header, master_sheet.title)
            _required_column(source_headers, header, source_sheet.title)
        source_only = sorted(header for header in source_headers if header not in master_headers)
        if source_only:
            raise ValueError(f"Headcount columns missing from '{master_sheet.title}': {source_only}")
        master_rows = _candidate_rows(master_sheet, master_header_row, master_headers, "master", master_headers)
        source_rows = _candidate_rows(source_sheet, source_header_row, source_headers, "source", master_headers)
        retained, superseded_rows = _retain_latest(master_rows + source_rows)
        duplicate_column = next((column for header, column in master_headers.items() if header.rstrip("?") in DUPLICATE_HEADERS), None)
        return HrRefreshPlan(
            master_sheet_name=master_sheet.title,
            master_header_row=master_header_row,
            master_headers=master_headers,
            retained_rows=[row["values"] for row in retained],
            existing_rows=len(master_rows),
            source_rows=len(source_rows),
            superseded_rows=superseded_rows,
            duplicate_column=duplicate_column,
        )
    finally:
        source_workbook.close()
        master_workbook.close()