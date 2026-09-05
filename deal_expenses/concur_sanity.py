"""In-scope YTD control checks for Concur workbooks."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from deal_expenses.models import RunRequest, ValidationResult
from deal_expenses.sources.base_concur import EXPENSE_HEADER, YEAR_HEADER, SourceAdapter, is_reporting_year, normalize_header


IN_SCOPE_SHEET_NAME = "In Scope CCs"
IN_SCOPE_NAME_HEADER = "CC Name"
IN_SCOPE_COST_CENTER_HEADER = "SAP LA CC"
CONCUR_SHEET_NAME = "Concur Report"
CONCUR_COST_CENTER_HEADER = "Org Unit 5 - Code"


def _decimal(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as error:
        raise ValueError(f"Expense amount {value!r} is not numeric.") from error


def _headers(worksheet) -> dict[str, int]:
    return {normalize_header(cell.value): index for index, cell in enumerate(worksheet[1], start=1) if normalize_header(cell.value)}


def _required_column(headers: dict[str, int], header: str, sheet_name: str) -> int:
    column = headers.get(header)
    if column is None:
        raise ValueError(f"Required header {header!r} was not found in worksheet {sheet_name!r}.")
    return column


def _cost_center_code(value: object) -> str:
    if value in (None, ""):
        return ""
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return str(value).strip()
    return str(number.quantize(Decimal("1"))) if number == number.to_integral_value() else format(number.normalize(), "f")


def load_concur_in_scope_cost_centers(master_path: Path) -> dict[str, str]:
    """Return the Concur cost-center code to name mapping from the master workbook."""
    workbook = load_workbook(master_path, read_only=True, data_only=True)
    try:
        if IN_SCOPE_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Worksheet {IN_SCOPE_SHEET_NAME!r} was not found.")
        worksheet = workbook[IN_SCOPE_SHEET_NAME]
        header_row = None
        name_column = code_column = None
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            headers = [normalize_header(value) for value in row]
            if IN_SCOPE_NAME_HEADER in headers and IN_SCOPE_COST_CENTER_HEADER in headers:
                header_row = row_number
                name_column = headers.index(IN_SCOPE_NAME_HEADER)
                code_column = headers.index(IN_SCOPE_COST_CENTER_HEADER)
                break
        if header_row is None or name_column is None or code_column is None:
            raise ValueError(f"Could not find the in-scope table on {IN_SCOPE_SHEET_NAME!r}.")

        cost_centers = {}
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            if all(value in (None, "") for value in row):
                break
            code = _cost_center_code(row[code_column])
            name = str(row[name_column]).strip() if row[name_column] not in (None, "") else ""
            if code and name:
                cost_centers[code] = name
    finally:
        workbook.close()
    if not cost_centers:
        raise ValueError(f"No SAP LA CC values were found on {IN_SCOPE_SHEET_NAME!r}.")
    return cost_centers


def _in_scope_total(worksheet, reporting_year: int, in_scope_cost_centers: set[str]) -> Decimal:
    headers = _headers(worksheet)
    year_column = _required_column(headers, YEAR_HEADER, worksheet.title) - 1
    expense_column = _required_column(headers, EXPENSE_HEADER, worksheet.title) - 1
    cost_center_column = _required_column(headers, CONCUR_COST_CENTER_HEADER, worksheet.title) - 1
    return sum(
        (
            _decimal(row[expense_column])
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            if is_reporting_year(row[year_column], reporting_year)
            and _cost_center_code(row[cost_center_column]) in in_scope_cost_centers
        ),
        Decimal("0"),
    )


class ConcurSanityChecker:
    """Compare current source YTD expenses against the prior master YTD total."""

    def __init__(self) -> None:
        self.in_scope_cost_centers: dict[str, str] = {}

    def validate(self, request: RunRequest, adapters: list[SourceAdapter]) -> ValidationResult:
        self.in_scope_cost_centers = load_concur_in_scope_cost_centers(request.master_path)
        in_scope_codes = set(self.in_scope_cost_centers)

        master_workbook = load_workbook(request.master_path, read_only=True, data_only=True)
        try:
            if CONCUR_SHEET_NAME not in master_workbook.sheetnames:
                raise ValueError(f"Worksheet {CONCUR_SHEET_NAME!r} was not found.")
            prior_total = _in_scope_total(master_workbook[CONCUR_SHEET_NAME], request.reporting_year, in_scope_codes)
        finally:
            master_workbook.close()

        current_total = Decimal("0")
        for adapter in adapters:
            workbook = load_workbook(request.source_paths[adapter.key], read_only=True, data_only=True)
            try:
                current_total += _in_scope_total(adapter.select_worksheet(workbook), request.reporting_year, in_scope_codes)
            finally:
                workbook.close()

        if current_total < prior_total:
            return ValidationResult(
                "Concur sanity check",
                False,
                f"Current in-scope YTD total ${current_total:,.2f} is lower than the prior total ${prior_total:,.2f}. Notify the team before continuing.",
            )
        return ValidationResult(
            "Concur sanity check",
            True,
            f"Performed: current in-scope YTD total ${current_total:,.2f} is at least the prior total ${prior_total:,.2f} across {len(in_scope_codes)} cost center(s).",
        )