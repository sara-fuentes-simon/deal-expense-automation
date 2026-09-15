from pathlib import Path

from openpyxl import Workbook

from deal_expenses.hr_refresh import build_hr_refresh_plan
from deal_expenses.models import UberRunRequest
from deal_expenses.sources import UberSourceAdapter
from deal_expenses.uber_pipeline import UberExpensePipeline
from deal_expenses.uber_validation import UberMasterWorkbookValidator, validate_source_compatibility


def create_master(path: Path) -> None:
    workbook = Workbook()
    uber = workbook.active
    uber.title = "Uber Report"
    headers = [
        "Request Date (UTC)",
        "First Name",
        "Last Name",
        "Full Name",
        "Employee ID",
        "Cost Center",
        "Cost Center",
        "Final Cost Center",
        "LOB",
        "In Scope?",
        "Transaction Amount USD",
        "Company Name",
    ]
    uber.append([])
    uber.append([])
    uber.append([])
    uber.append([])
    uber.append([None, None, None, None, None, "Using Employee ID", "Using Employee Name"])
    uber.append(headers)
    uber.append(["2026-01-01", "Ada", "Lovelace", '=CONCATENATE(B7," ",C7)', "1", "CC-1", None, "=IF(F7<>\"\",F7,G7)", "=H7", "=\"Yes\"", 10, "=H7"])
    hr = workbook.create_sheet("HR HC Combined Jun")
    hr.append(["Effective Date", "Network ID", "Employee ID", "Name", "Cost Center Name", "DUPLICATE?"])
    hr.append(["2026-01-01", "1", "1", "Ada Lovelace", "CC-1", '=COUNTIF(B:B,B2)>1'])
    workbook.save(path)


def create_source(path: Path, extra_header: str | None = None) -> None:
    workbook = Workbook()
    source = workbook.active
    source.title = "U4B"
    headers = ["Request Date (UTC)", "First Name", "Last Name", "Employee ID", "Transaction Amount USD"]
    if extra_header:
        headers.append(extra_header)
    source.append(headers)
    source.append(["2026-01-01", "Ada", "Lovelace", "1", 10] + (["extra"] if extra_header else []))
    workbook.save(path)


def create_headcount(path: Path) -> None:
    workbook = Workbook()
    source = workbook.active
    source.title = "To be Copy Pasted from HR"
    source.append(["Headcount export"])
    source.append([])
    source.append([])
    source.append(["Effective Date", "Network ID", "Employee ID", "Name", "Cost Center Name"])
    source.append(["2026-02-01", "1", "1", "Ada Lovelace", "CC-2"])
    workbook.save(path)


def test_uber_adapter_selects_only_eligible_worksheet(tmp_path: Path):
    source_path = tmp_path / "uber.xlsx"
    create_source(source_path)

    adapter = UberSourceAdapter()

    assert adapter.validate(source_path)[0].passed
    assert adapter.summarize(source_path, 2026).row_count == 1


def test_uber_mapping_rejects_source_only_header(tmp_path: Path):
    master_path = tmp_path / "master.xlsx"
    source_path = tmp_path / "uber.xlsx"
    create_master(master_path)
    create_source(source_path, "Unexpected Field")

    result = validate_source_compatibility(master_path, source_path, UberSourceAdapter())

    assert not result.passed
    assert "Source-only columns" in result.message


def test_uber_master_requires_final_cost_center_formula_template(tmp_path: Path):
    master_path = tmp_path / "master.xlsx"
    create_master(master_path)

    from openpyxl import load_workbook

    workbook = load_workbook(master_path)
    workbook["Uber Report"].cell(7, 8).value = "CC-1"
    workbook.save(master_path)
    workbook.close()

    result = UberMasterWorkbookValidator().validate(master_path)

    assert not result[0].passed
    assert "Final Cost Center" in result[0].message


def test_headcount_plan_keeps_the_latest_network_id_record(tmp_path: Path):
    master_path = tmp_path / "master.xlsx"
    headcount_path = tmp_path / "headcount.xlsx"
    create_master(master_path)
    create_headcount(headcount_path)

    plan = build_hr_refresh_plan(master_path, headcount_path)

    assert plan.source_rows == 1
    assert plan.superseded_rows == 1
    assert plan.retained_rows[0][4] == "CC-2"


class StubUberWriter:
    def write(self, request: UberRunRequest) -> dict[str, int | float]:
        return {
            "imported_rows": 1,
            "unassigned_final_cost_center_rows": 1,
            "unassigned_final_cost_center_percent": 100.0,
        }


def test_uber_pipeline_reports_unassigned_final_cost_centers_without_stopping(tmp_path: Path):
    master_path = tmp_path / "master.xlsx"
    headcount_path = tmp_path / "headcount.xlsx"
    source_path = tmp_path / "uber.xlsx"
    create_master(master_path)
    create_headcount(headcount_path)
    create_source(source_path)
    pipeline = UberExpensePipeline(UberSourceAdapter(), UberMasterWorkbookValidator(), StubUberWriter())

    events = list(pipeline.run(UberRunRequest(master_path, headcount_path, source_path, 2026)))

    assert pipeline.result.success
    assert pipeline.metrics["unassigned_final_cost_center_rows"] == 1
    assert events[-1].step == "Complete"
